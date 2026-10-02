"""Build results/GEMV_Charts_XRT.xlsx -- ONE workbook with every result of the measurement
windows (native XRT host, ert=false, quiet server, 2026-10-01 and -02), laid out for the charts
the Claude Excel extension draws into it.

    python scripts/analysis/make_xrt_charts_xlsx.py                       # the chart workbook
    python scripts/analysis/make_xrt_charts_xlsx.py --out results/GEMV_Charts_XRT_values_<date>.xlsx

ONE-SHOT: it REFUSES to overwrite an existing workbook. Once the extension has drawn charts into
GEMV_Charts_XRT.xlsx the file is the user's -- a rebuild deletes every chart, and openpyxl
cannot carry charts through a load and save either. NEW DATA is therefore PASTED: --out builds
a VALUES workbook with exactly the same sheets and cell layout from the current CSVs, and the
user copies each data sheet's block into the chart workbook with Paste Special > Values; the
charts address cells, so they follow. Columns are only ever appended (make_window_csv.py), so a
pasted block never moves a column a chart uses.

Inputs : results/GEMV_Qwen_XRT.csv, GEMV_Qwen_Accuracy.csv, GEMV_UC1_XRT.csv (make_window_csv.py)
Sheets : Qwen Charts, UC1 Charts  empty, for the charts
         Qwen XRT        UC2 per build, the CSV as is (sorted by HBM channels; row 23 pending)
         Qwen Twins      own / shared / broadcast vector for the three twin sets: one small
                         block per metric, each block one contiguous chart range
         Qwen Scaling    4x4 engines 1..7 by vector feed: tokens/s, mJ/token, clock
         Qwen Accuracy   the realistic layer's error by sparsity, the CSV as is
         UC1 XRT         UC1 per build: the OpenCL 'Workload per matrix' columns + the window's
         Notes           what was measured, under which conditions, what each sheet holds
"""

import argparse
import csv
import io
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))  # repository root; this file is in scripts/analysis/
RESULTS = os.path.join(ROOT, "results")
sys.path.insert(0, HERE)
import make_chart_xlsx as mcx                   # noqa: E402  the same sheet writer and styling

OUT = os.path.join(RESULTS, "GEMV_Charts_XRT.xlsx")
QWEN, ACC, UC1 = "GEMV_Qwen_XRT.csv", "GEMV_Qwen_Accuracy.csv", "GEMV_UC1_XRT.csv"
TWINS = [("3 x 4x4", "3x4x4"), ("3 x 8x4", "3x8x4"), ("2 x 16x3", "2x16x3")]
FEEDS = [("own vector", ""), ("shared vector", "_shared"), ("broadcast vector", "_bcast")]
# (block title, column of GEMV_Qwen_XRT.csv, scale, number format)
TWIN_METRICS = [("Time per token (ms)", "token_us", 1e-3, "0.00"),
                ("Tokens per second", "tokens_per_s", 1, "0.0"),
                ("Energy per token (mJ)", "energy_token_mJ", 1, "0.0"),
                ("Board idle power (W)", "board_idle_W", 1, "0.00"),
                ("HBM channels", "hbm_channels", 1, "0"),
                ("Clock (MHz)", "clock_mhz", 1, "0")]
BOLD = Font(bold=True, size=12)


def fail(msg):
    raise SystemExit("make_xrt_charts_xlsx: " + msg + " -- nothing written")


def read_csv(name):
    path = os.path.join(RESULTS, name)
    if not os.path.exists(path):
        fail("no %s -- run make_window_csv.py first" % os.path.relpath(path, ROOT))
    with io.open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def header(ws, row, names):
    for i, name in enumerate(names, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.fill, c.font, c.border = mcx.H_FILL, mcx.H_FONT, mcx.BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def value(ws, row, col, v, fmt=None):
    c = ws.cell(row=row, column=col, value=v)
    c.border = mcx.BOX
    if fmt and isinstance(v, (int, float)):
        c.number_format = fmt
    return c


def by_arch(qwen):
    out = {}
    for r in qwen:
        if r["status"] != "measured":
            continue
        # source lists every run, " + "-joined: any one names the build
        m = re.search(r"workload_xrt_qwen_(.+)_\d+MHz\.csv$", r["source"].split(" + ")[-1])
        if not m:
            fail("cannot tell the build of row %s" % r["config_label"])
        out[m.group(1)] = r
    return out


def twins_sheet(wb, arch):
    ws = wb.create_sheet("Qwen Twins")
    for k, (title, col, scale, fmt) in enumerate(TWIN_METRICS):
        top = 6 * k + 1                              # title, header, 3 rows, a blank row
        ws.cell(row=top, column=1, value=title).font = BOLD
        header(ws, top + 1, ["twin set"] + [f for f, _ in FEEDS])
        for i, (label, base) in enumerate(TWINS):
            value(ws, top + 2 + i, 1, label)
            for j, (_f, suffix) in enumerate(FEEDS):
                name = base + suffix
                if name not in arch:
                    fail("twin %s not measured" % name)
                v = float(arch[name][col]) * scale
                value(ws, top + 2 + i, 2 + j, int(round(v)) if fmt == "0" else round(v, 3), fmt)
    for c, w in zip("ABCD", (22, 18, 18, 18)):
        ws.column_dimensions[c].width = w


def scaling_sheet(wb, arch):
    ws = wb.create_sheet("Qwen Scaling")
    ws.cell(row=1, column=1, value="4x4 engines on the Qwen test, by vector feed "
                                   "(a blank = no such build)").font = BOLD
    cols = [("tokens/s", "tokens_per_s", "0.0"), ("mJ/token", "energy_token_mJ", "0.0"),
            ("clock MHz", "clock_mhz", "0")]
    names = ["4x4 engines"] + ["%s: %s" % (f, unit) for unit, _c, _m in cols for f, _s in FEEDS]
    header(ws, 2, names)
    for n in range(1, 8):
        row = 2 + n
        value(ws, row, 1, n)
        for u, (_unit, col, fmt) in enumerate(cols):
            for j, (_f, suffix) in enumerate(FEEDS):
                name = ("4x4" if n == 1 else "%dx4x4" % n) + suffix
                if n == 1 and suffix:
                    continue                          # one engine has its own vector
                if name in arch:
                    v = float(arch[name][col])
                    value(ws, row, 2 + 3 * u + j, int(round(v)) if fmt == "0" else v, fmt)
    ws.column_dimensions["A"].width = 12
    for i in range(2, 11):
        ws.column_dimensions[chr(64 + i)].width = 17


def notes_sheet(wb, qwen, uc1):
    ws = wb.create_sheet("Notes")
    window = sorted(set(w for r in qwen + uc1 if r.get("window") for w in r["window"].split(" + ")))
    meas = [r for r in qwen if r["status"] == "measured"]
    nruns = sorted(set(r.get("runs", "1") for r in meas + uc1))
    t_hr = max(100.0 * float(r["runs_token_ms_halfrange"]) / float(r["token_ms_" + k])
               for r in meas for k in ("single", "multi", "shared", "bcast")
               if r.get("token_ms_" + k) and r.get("runs_token_ms_halfrange"))
    dyn = max((float(r["runs_dynamic_W_max"]) - float(r["runs_dynamic_W_min"]), r["config_label"])
              for r in meas if r.get("runs_dynamic_W_max"))
    e_hr = max((100.0 * float(r["runs_energy_token_mJ_halfrange"]) / float(r["energy_token_mJ"]),
                r["config_label"]) for r in meas if r.get("runs_energy_token_mJ_halfrange"))
    capped = ["%s %.1f %% (%.0f M beats/s at %s MHz)" % (
        r["config_label"].split(" (")[0], float(r["efficiency_pct"]),
        float(r["efficiency_pct"]) / 100.0 * float(r["clock_mhz"]), r["clock_mhz"])
        for r in meas if float(r["efficiency_pct"]) < 90.0]
    pending = [r["config_label"] for r in qwen if r["status"] != "measured"]
    lines = [
        ("GEMV_Charts_XRT.xlsx -- the measurement windows of 2026-10-01 and 2026-10-02", True),
        ("Runs: every build measured %s times on a quiet server (window folders results/%s/); the "
         "charts show the MEAN of the runs, and the runs_* columns at the right of 'Qwen XRT' and "
         "'UC1 XRT' hold each quantity's min / max / half-range over them, for error bars"
         % (" or ".join(nruns), ", results/".join(window)), False),
        ("Conditions: 1 Oct -- the server alone after a reboot; 2 Oct -- the professor's isolated "
         "window; both: no Vivado / v++ / HLS of anyone (30 s load logs results/window_load_*.txt), "
         "1-minute load ~3.4 during the runs = their own processes", False),
        ("Host: native XRT C++ (host_workload_xrt.cpp), ert=false = host-side scheduling (the "
         "driver starts the CUs, not the card's ERT); XRT 2.13.0 (2022.1); Alveo U280", False),
        ("Measured: wall time = a 60 s soak of whole passes / tokens; board power from the card's "
         "sensors; static energy = idle power x time, dynamic = (load - idle) x time", False),
        ("Derived: engine time = the plan's cycles / clock / measured steady-state efficiency -- "
         "the build's GEMV_Configurations.csv value (13 builds), or for a shared or broadcast "
         "build its OWN MIXED shape-sweep occupancy (min of the 11 shapes); overhead = wall - "
         "engine (launches, completion, lockstep)", False),
        ("Finding: %s -- a hard cap on the weight stream (half of the 450 MHz HBM port clock), "
         "the same on every shape and whatever the kernel clock; every other build streams at "
         "~99 %% of its clock" % ("; ".join(capped) if capped else "no build below 90 %"), False),
        ("Reproducibility between the runs: time per token / pass within %.1f %% (half-range); "
         "the Qwen test's dynamic power moved by up to %.1f W (%s), its energy per token by up "
         "to %.1f %% half-range (%s) -- random data, a new seed every run"
         % (t_hr, dyn[0], dyn[1], e_hr[0], e_hr[1]), False),
        ("UC2 Qwen3.5-35B-A3B step 1: per token 40 layers, each a 9216 x 2048 MIXED stacked GEMV "
         "(9 experts x gate+up, 2304 rows at each of 2:4 / 2:8 / 2:16 / 2:32), split over the "
         "engines in lockstep, a new vector every layer; weights N(0,1) magnitude-pruned, vector "
         "N(0,1)", False),
        ("UC1 multi-user: 132 matrices, every matrix its own calculation and input vector -- the "
         "same columns and rows as the OpenCL sheet 'Workload per matrix' (2026-09-28)", False),
        ("Accuracy: the gates' realistic layer, |card - exact| / rms(exact); the card matches "
         "'each summand toward zero to 2^-7, fixed-point sum, bf16 truncation' on 100 % of rows. "
         "Twins (own / shared / broadcast) were gated on the same data and returned the same bits",
         False),
        ("Chart colours: single engine #2a78d6, multi-tenant own vector #eb6834, shared vector "
         "#1baf7a, broadcast vector #eda100; stacked: the kind colour (engine / static) under its "
         "light tint (overhead / dynamic)", False),
        (("Pending: %s -- blank rows, already inside every chart range" % ", ".join(pending))
         if pending else "Pending: none -- every build measured", False),
        ("Caveat: the OpenCL UC1 run (2026-09-28) differs in host API, scheduler (ERT) and server "
         "load at once -- not an equal-conditions comparison", False),
        ("Updating: NEVER rebuild this file (the builder refuses -- a rebuild deletes the charts); "
         "build a values workbook (make_xrt_charts_xlsx.py --out ...) and paste each data sheet's "
         "block over this one with Paste Special > Values", False),
        ("Sources: results/GEMV_Qwen_XRT.csv, GEMV_Qwen_Accuracy.csv, GEMV_UC1_XRT.csv "
         "(scripts/analysis/make_window_csv.py); this file: scripts/analysis/"
         "make_xrt_charts_xlsx.py", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = BOLD
    ws.column_dimensions["A"].width = 140


def main():
    ap = argparse.ArgumentParser(description="the XRT chart workbook, or a values workbook to "
                                             "paste into it")
    ap.add_argument("--out", default=OUT, help="the workbook to write (default %s); it must not "
                                               "exist yet" % os.path.relpath(OUT, ROOT))
    a = ap.parse_args()
    out = os.path.abspath(a.out)
    if os.path.exists(out):
        fail("%s exists -- it may hold charts, and a rebuild deletes them; paste new data into "
             "it instead (or delete it yourself if it has no charts)" % os.path.relpath(out, ROOT))
    lock = os.path.join(os.path.dirname(out), "~$" + os.path.basename(out))
    if os.path.exists(lock):
        fail("Excel has %s open" % os.path.basename(out))
    qwen, uc1 = read_csv(QWEN), read_csv(UC1)
    read_csv(ACC)
    arch = by_arch(qwen)

    wb = Workbook()
    wb.remove(wb.active)
    for title, text in (("Qwen Charts", "UC2 -- Qwen3.5-35B-A3B step 1 on the U280, native XRT "
                                        "host (ert=false), quiet server, 2026-10-01 / -02"),
                        ("UC1 Charts", "UC1 -- multi-user, every matrix its own calculation and "
                                       "vector, native XRT host (ert=false), 2026-10-01 / -02")):
        ws = wb.create_sheet(title)
        ws.cell(row=1, column=1, value=text).font = BOLD
    made = [mcx.write_sheet(wb, QWEN, "Qwen XRT")]
    twins_sheet(wb, arch)
    scaling_sheet(wb, arch)
    made.append(mcx.write_sheet(wb, ACC, "Qwen Accuracy"))
    made.append(mcx.write_sheet(wb, UC1, "UC1 XRT"))
    notes_sheet(wb, qwen, uc1)
    if not all(made):
        fail("an empty table")
    wb.save(out)
    print("wrote %s: %s" % (os.path.relpath(out, ROOT), ", ".join(wb.sheetnames)))


if __name__ == "__main__":
    main()
