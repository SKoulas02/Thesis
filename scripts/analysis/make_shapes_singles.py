"""The 8x8 shape charts, redone for every single-engine family build: latency per matrix,
silicon rate, energy per element and energy per matrix, 11 shapes x {2:4, 2:8, 2:16, 2:32,
MIXED}, each engine at its own closed clock. No card time: the family campaign (2026-09-13/14)
already ran exactly this workload -- run_shapes.py batched sweeps + one 60 s soak per mode.

    python scripts/analysis/make_shapes_singles.py             # the CSV (always rewritten)
    python scripts/analysis/make_shapes_singles.py --xlsx      # + the chart workbook, ONCE

Inputs : results/family_measurements/<tag>/shapes_<tag>_<MHz>MHz_<mode>.csv  (check_family_measure)
         results/family_measurements/<tag>/power_<tag>_<MHz>MHz.csv           one soak per mode
Outputs: results/GEMV_Shapes_Singles.csv      LONG record, 330 rows (6 builds x 5 modes x 11)
         results/GEMV_Charts_Shapes_Singles.xlsx  chart-ready blocks + empty chart sheets; built
                                               ONCE (refuses to overwrite: it will hold charts)

LATENCY = the family campaign's steady-state estimator: (mean of the three batched runs - the
POOLED launch overhead) / R, the overhead pooled over the build's five sweeps by inverse
variance (check_family_measure.steady -- the estimator every cross-family chart uses). The
MIXED-only sheet of 2026-09-18 used each sweep's own fit instead; the two differ by <= 0.68 %
there, and the largest difference anywhere is printed. PADDING IS INCLUDED (the real time to
finish the matrix; row granularity = lanes for one sparsity, 4 x lanes for MIXED -- checked).

SILICON RATE = GFLOPS actually computed (2 FLOP x 2 MAC x lanes per weight beat, padded beats
included) -- flat across shapes and modes. DELIVERED = 2 x M x N / latency (dense-equivalent).
ENERGY = power x time with the mode's OWN soak: static = idle x latency, dynamic = (load - idle)
x latency; per element divides by the TRUE M x N. Power was soaked at V=1024 and applied to
every shape, as for the 8x8 charts. Dynamic power per mode is CONFOUNDED with soak order (2:4
always first) -- do not read sparsity into the small dynamic differences.

LONG blocks put the outer category label (the shape) on the FIRST row of each 5-row group only
-- Excel starts a group at every non-blank outer cell.
"""

import argparse
import csv
import io
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))  # repository root; this file is in scripts/analysis/
RES = os.path.join(ROOT, "results")
sys.path.insert(0, HERE)
import check_family_measure as cfm             # noqa: E402  loader, CONFIGS, ORDER, steady()

OUT_CSV = os.path.join(RES, "GEMV_Shapes_Singles.csv")
OUT_XLSX = os.path.join(RES, "GEMV_Charts_Shapes_Singles.xlsx")
MODES = [lab for _f, lab, _c in cfm.LABS]       # 2:4 2:8 2:16 2:32 MIXED
LAT_BLOCK, LONG_BLOCK = 15, 60                   # rows per block in the WIDE / LONG sheets


def fail(msgs):
    print("CHECKS FAILED -- nothing written:")
    for m in msgs:
        print("   ", m)
    raise SystemExit(1)


def collect():
    rows, problems, ref, worst = [], [], None, (0.0, "")
    for tag in cfm.ORDER:
        C, B, clk = cfm.CONFIGS[tag]
        L = C * B
        o = cfm.load(tag)
        if o is None:
            fail(["missing family measurements for %s" % tag])
        o["L"] = L
        shapes_here = None
        for mode in MODES:
            S = o["shapes"][mode]
            pw = [r for r in o["power"] if r["label"].split()[-1] == mode]
            if len(S) != 11 or len(pw) != 1:
                problems.append("%s %s: %d shapes, %d power rows (want 11, 1)"
                                % (tag, mode, len(S), len(pw)))
                continue
            load_w, idle_w = float(pw[0]["board_load_w"]), float(pw[0]["board_idle_w"])
            if not idle_w < load_w or int(float(pw[0]["clock_mhz"])) != clk:
                problems.append("%s %s: soak load %.2f / idle %.2f W at %s MHz"
                                % (tag, mode, load_w, idle_w, pw[0]["clock_mhz"]))
            pool, steady = cfm.steady(o, mode)
            grain = 4 * L if mode == "MIXED" else L
            got = []
            for r, st in zip(S, steady):
                M, N = int(r["M_rows"]), int(r["N_cols"])
                if r["sparsity"] != mode or int(float(r["clock_mhz"])) != clk:
                    problems.append("%s %s %s: file says %s @ %s" % (tag, mode, r["shape"],
                                                                    r["sparsity"], r["clock_mhz"]))
                pad = int(r["padding_rows"] or 0)
                if pad != int(math.ceil(float(M) / grain)) * grain - M:
                    problems.append("%s %s %s: padding %d breaks the %d-row grain"
                                    % (tag, mode, r["shape"], pad, grain))
                per = float(r["ideal_per_matrix_us"]) / st["occ"]       # pooled latency
                own = float(r["latency_per_matrix_us"])                  # the sweep's own fit
                dev = 100.0 * abs(own / per - 1.0)
                worst = max(worst, (dev, "%s %s %s" % (tag, mode, r["shape"])))
                el = float(M * N)
                st_uj, dy_uj = idle_w * per, (load_w - idle_w) * per
                got.append((r["shape"], M, N))
                rows.append(dict(
                    build=tag, clock_mhz=clk, lanes=L, sparsity=mode, shape=r["shape"],
                    dimensions="%dx%d" % (M, N), M_rows=M, N_cols=N, matrix_elements=M * N,
                    padding_rows=pad, latency_us=round(per, 4),
                    latency_own_fit_us=round(own, 4), own_vs_pooled_pct=round(dev, 3),
                    pooled_overhead_us=round(pool, 1), batch_R=int(r["batch_R"]),
                    occupancy=round(st["occ"], 4), gflops_silicon=round(st["gfa"], 3),
                    gflops_delivered=round(st["gfe"], 3), board_load_W=load_w,
                    board_idle_W=idle_w, board_dynamic_W=round(load_w - idle_w, 3),
                    energy_static_uJ=round(st_uj, 4), energy_dynamic_uJ=round(dy_uj, 4),
                    energy_total_uJ=round(st_uj + dy_uj, 4),
                    static_pJ_per_element=round(st_uj / el * 1e6, 3),
                    dynamic_pJ_per_element=round(dy_uj / el * 1e6, 3),
                    total_pJ_per_element=round((st_uj + dy_uj) / el * 1e6, 3),
                    static_pct=round(100.0 * st_uj / (st_uj + dy_uj), 2),
                    data_source="family_measurements/%s: shapes_%s_%dMHz_<mode>.csv (pooled "
                                "overhead) x power_%s_%dMHz.csv (the mode's 60 s soak)"
                                % (tag, tag, clk, tag, clk)))
            shapes_here = shapes_here or got
            if got != shapes_here:
                problems.append("%s %s: shape list differs from its other sweeps" % (tag, mode))
        if ref is None:
            ref = shapes_here
        elif shapes_here != ref:
            problems.append("%s: its 11 shapes differ from %s's" % (tag, cfm.ORDER[0]))
    if problems:
        fail(problems)
    return rows, ref, worst


def write_csv(rows):
    with io.open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print("wrote %s (%d rows, %d columns)" % (os.path.relpath(OUT_CSV, ROOT), len(rows),
                                             len(rows[0])))


def write_xlsx(rows, ref):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    import make_chart_xlsx as mcx

    if os.path.exists(OUT_XLSX):
        raise SystemExit("%s exists -- it may hold charts and a rebuild deletes them; nothing "
                         "written" % os.path.relpath(OUT_XLSX, ROOT))
    if os.path.exists(os.path.join(RES, "~$" + os.path.basename(OUT_XLSX))):
        raise SystemExit("Excel has %s open -- nothing written" % os.path.basename(OUT_XLSX))
    by = dict(((r["build"], r["sparsity"], r["shape"]), r) for r in rows)
    bold = Font(bold=True, size=12)
    wb = Workbook()
    wb.remove(wb.active)
    for tag in cfm.ORDER:
        ws = wb.create_sheet("Charts %s" % tag)
        ws.cell(row=1, column=1, value="%s @ %d MHz -- the 8x8 shape charts for this engine "
                                       "(family campaign data)" % (tag, cfm.CONFIGS[tag][2])
                ).font = bold

    def cell(ws, r, c, v, head=False, fmt=None):
        x = ws.cell(row=r, column=c, value=v)
        x.border = mcx.BOX
        if head:
            x.fill, x.font = mcx.H_FILL, mcx.H_FONT
        elif fmt and isinstance(v, float):
            x.number_format = fmt
        return x

    for title, key, fmt in (("Latency", "latency_us", "0.000"),
                            ("Silicon Rate", "gflops_silicon", "0.0")):
        ws = wb.create_sheet(title)
        for b, tag in enumerate(cfm.ORDER):
            top = 1 + LAT_BLOCK * b
            ws.cell(row=top, column=1, value="%s @ %d MHz" % (tag, cfm.CONFIGS[tag][2])
                    ).font = bold
            for j, h in enumerate(["dimensions"] + MODES):
                cell(ws, top + 1, 1 + j, h, head=True)
            for i, (shape, M, N) in enumerate(ref):
                cell(ws, top + 2 + i, 1, "%dx%d" % (M, N))
                for j, mode in enumerate(MODES):
                    cell(ws, top + 2 + i, 2 + j, by[(tag, mode, shape)][key], fmt=fmt)
        ws.column_dimensions["A"].width = 14
    for title, keys, fmt in (("Energy per Element", ("static_pJ_per_element",
                                                     "dynamic_pJ_per_element"), "0.0"),
                             ("Energy per Matrix", ("energy_static_uJ",
                                                    "energy_dynamic_uJ"), "0.00")):
        ws = wb.create_sheet(title)
        unit = "pJ/element" if "Element" in title else "uJ"
        for b, tag in enumerate(cfm.ORDER):
            top = 1 + LONG_BLOCK * b
            ws.cell(row=top, column=1, value="%s @ %d MHz" % (tag, cfm.CONFIGS[tag][2])
                    ).font = bold
            for j, h in enumerate(["shape", "sparsity", "Static (%s)" % unit,
                                   "Dynamic (%s)" % unit]):
                cell(ws, top + 1, 1 + j, h, head=True)
            r = top + 2
            for shape, M, N in ref:
                for j, mode in enumerate(MODES):
                    x = by[(tag, mode, shape)]
                    cell(ws, r, 1, "%dx%d" % (M, N) if j == 0 else None)   # FIRST row only
                    cell(ws, r, 2, mode)
                    cell(ws, r, 3, x[keys[0]], fmt=fmt)
                    cell(ws, r, 4, x[keys[1]], fmt=fmt)
                    r += 1
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = 20
    mcx.write_sheet(wb, OUT_CSV, "Data")
    ws = wb.create_sheet("Notes")
    for i, t in enumerate([
            "GEMV_Charts_Shapes_Singles.xlsx -- the 8x8 shape charts for the six single-engine "
            "family builds, each at its own clock (4x4 400, 8x4 375, 16x3 350, 8x8 325, 4x24 "
            "300, 4x32 250 MHz)",
            "Data: the family campaign of 2026-09-13/14 (run_shapes.py batched sweeps, OpenCL "
            "host); latency is compute-only (launch overhead removed: pooled per-build fit), "
            "so the host does not enter it",
            "Padding included (real time to finish the matrix). Silicon rate = GFLOPS actually "
            "computed; delivered = 2 x M x N / latency",
            "Energy = the mode's own 60 s soak (V=1024) x latency; static = idle x time, dynamic "
            "= (load - idle) x time; per element = / true M x N",
            "Caveat: dynamic power per mode is confounded with soak order (2:4 always first) -- "
            "do not read sparsity into its small differences",
            "Never rebuild this file once it holds charts (the builder refuses). Source: "
            "results/GEMV_Shapes_Singles.csv, scripts/analysis/make_shapes_singles.py"], start=1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 140
    wb.save(OUT_XLSX)
    print("wrote %s: %s" % (os.path.relpath(OUT_XLSX, ROOT), ", ".join(wb.sheetnames)))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsx", action="store_true", help="also build the chart workbook (once)")
    a = ap.parse_args()
    rows, ref, worst = collect()
    write_csv(rows)
    print("pooled vs each sweep's own overhead fit: worst %.2f %% (%s)" % worst)
    print("\n%-6s %-6s %9s %9s %9s %9s %9s   silicon GFLOPS (min..max over 55)  static %%"
          % ("build", "MHz", "2:4 us", "2:8", "2:16", "2:32", "MIXED"))
    for tag in cfm.ORDER:
        mine = [r for r in rows if r["build"] == tag]
        g = [r["gflops_silicon"] for r in mine]
        lat = dict((r["sparsity"], r["latency_us"]) for r in mine if r["dimensions"] == "4096x4096")
        print("%-6s %-6d %9.2f %9.2f %9.2f %9.2f %9.2f   %7.1f .. %7.1f              %5.1f-%4.1f"
              % (tag, cfm.CONFIGS[tag][2], lat["2:4"], lat["2:8"], lat["2:16"], lat["2:32"],
                 lat["MIXED"], min(g), max(g), min(r["static_pct"] for r in mine),
                 max(r["static_pct"] for r in mine)))
    print("(latencies shown for 4096x4096)")
    if a.xlsx:
        write_xlsx(rows, ref)


if __name__ == "__main__":
    main()
